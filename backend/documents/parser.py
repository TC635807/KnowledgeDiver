"""
文档解析工具模块。

支持将 .txt / .md / .pdf / .docx 文件解析为纯文本，
文件大小上限 20 MB，自动处理编码格式和文件损坏检测。
"""

from __future__ import annotations

import logging
from typing import Optional

logger = logging.getLogger(__name__)

_MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB 文件大小上限

_UNSUPPORTED_DOC_HINT = (
    "旧版 .doc 格式（Microsoft Word 97-2003）不支持直接解析。"
    "请用 Word / WPS 将文件另存为 .docx 格式后重新上传。"
)


class ParseError(Exception):
    """文件解析错误。"""


def parse_file(filename: str, content: bytes) -> str:
    """将上传的文件解析为纯文本。

    根据文件扩展名选择解析方式，文件大小超过 20 MB 直接拒绝。
    未知扩展名尝试按文本解码。

    Args:
        filename: 文件名（用于判断扩展名）
        content: 文件二进制内容

    Returns:
        提取的纯文本

    Raises:
        ParseError: 不支持格式、文件过大或解析失败
    """
    if len(content) > _MAX_FILE_SIZE:
        raise ParseError(f"文件过大（{len(content) / 1024 / 1024:.1f} MB），上限 20 MB")

    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if ext in ("txt", "md", "markdown", "text"):
        return _parse_text(content)
    elif ext == "pdf":
        return _parse_pdf(content)
    elif ext == "docx":
        return _parse_docx(content)
    elif ext == "doc":
        raise ParseError(_UNSUPPORTED_DOC_HINT)
    else:
        # 未知扩展名回退文本解码
        try:
            return _parse_text(content)
        except UnicodeDecodeError:
            raise ParseError(f"不支持的文件格式: .{ext}（支持 txt, md, pdf, docx）")


def _parse_text(content: bytes) -> str:
    """解码纯文本文件，按顺序尝试 UTF-8 → GBK → GB2312 → latin-1。"""
    for encoding in ("utf-8", "gbk", "gb2312", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("无法解码文件，请确认文件编码为 UTF-8 或 GBK")


def _parse_pdf(content: bytes) -> str:
    """使用 pdfplumber 解析 PDF 文件，逐页提取文本。"""
    from io import BytesIO
    import pdfplumber

    try:
        doc = pdfplumber.open(BytesIO(content))
    except Exception as e:
        raise ParseError(f"PDF 文件损坏或无法读取: {e}")

    pages: list[str] = []
    try:
        for page in doc.pages:
            text = page.extract_text()
            if text and text.strip():
                pages.append(text)
    finally:
        doc.close()

    if not pages:
        raise ParseError("PDF 文件中没有可提取的文字（可能是扫描图片版，不支持 OCR）")

    return "\n\n".join(pages)


def _parse_docx(content: bytes) -> str:
    """使用 python-docx 解析 DOCX 文件，提取段落和表格文本。"""
    from io import BytesIO
    from docx import Document

    try:
        doc = Document(BytesIO(content))
    except Exception as e:
        raise ParseError(f"DOCX 文件损坏或无法读取: {e}")

    paragraphs: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            paragraphs.append(text)

    # 提取表格文本
    for table in doc.tables:
        for row in table.rows:
            row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_texts:
                paragraphs.append(" | ".join(row_texts))

    if not paragraphs:
        raise ParseError("DOCX 文件中没有可提取的文字内容")

    return "\n\n".join(paragraphs)
